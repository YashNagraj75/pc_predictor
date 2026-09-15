"""Build the progress report as .docx so Drive converts it to a native Google Doc."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import re

ACC = RGBColor(0x1B, 0x6C, 0xA8)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()
for s in ("Normal",):
    st = doc.styles[s]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin = sec.bottom_margin = Inches(0.85)


def rich(p, text, base_size=10.5, color=None):
    """Render a mini-markup string: **bold**, *italic*, ~sub~, ^sup^."""
    tokens = re.split(r"(\*\*.+?\*\*|\*.+?\*|\^.+?\^|~.+?~)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            r = p.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith("^") and tok.endswith("^"):
            r = p.add_run(tok[1:-1]); r.font.superscript = True
        elif tok.startswith("~") and tok.endswith("~"):
            r = p.add_run(tok[1:-1]); r.font.subscript = True
        else:
            r = p.add_run(tok)
        r.font.size = Pt(base_size)
        if color is not None:
            r.font.color.rgb = color


def para(text, size=10.5, color=None, space_after=7, justify=True, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    rich(p, text, size, color)
    return p


def heading(text, level=1):
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(13)
    h.paragraph_format.space_after = Pt(5)
    r = h.add_run(text)
    r.font.color.rgb = ACC
    r.font.size = Pt(13 if level == 1 else 11.5)
    r.font.name = "Calibri"
    return h


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    rich(p, text)
    return p


def table(rows, widths, header=True):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cells[ci].width = Inches(widths[ci])
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.size = Pt(9)
            if header and ri == 0:
                r.bold = True
    return t


# ------------------------------------------------------------------ title
tp = doc.add_paragraph()
tp.paragraph_format.space_after = Pt(2)
tr = tp.add_run("Predictive Coding Predictors for Joint-Embedding World Models")
tr.bold = True
tr.font.size = Pt(18)

para("Progress report and request for feedback  |  Yash Nagraj  |  August 2026",
     size=10, color=GREY, space_after=12, justify=False)

para("**Where this stands.** The project began as a proposal to train the "
     "predictor of an I-JEPA-style joint-embedding world model with predictive "
     "coding (PC) instead of backpropagation (BP). The original motivation was "
     "that PC might supply a better gradient signal. I now think that motivation "
     "was wrong, and I have replaced it with a different one. This report covers "
     "what I built, what I measured, what changed my mind, and three questions I "
     "would like your opinion on. Everything numerical below is from code I wrote "
     "and ran; nothing is quoted from a paper.")

# ------------------------------------------------------------------ 1
heading("1. The reframing: PC does not give better gradients")
para("Millidge, Tschantz and Buckley (2020) show that PC approximates BP along "
     "arbitrary computation graphs. So \u201cPC gives a better gradient than "
     "BP\u201d is close to self-defeating: to the extent PC works, it works "
     "*because* it recovers the BP gradient. I verified this myself rather than "
     "take it on faith \u2014 I implemented PC over an explicit computation graph "
     "in JAX and compared it against autodiff on three topologies (a plain chain, "
     "a residual network with folded skips, and one with explicit skips). The PC "
     "and BP weight gradients agree to **cosine similarity 1.000000** on all "
     "three. On a single branch the PC gradient is **\u221211.370010** against "
     "BP\u2019s **\u221211.370533**.")
para("Two things fell out of that exercise that I did not expect, and that I "
     "think are the more useful results:")
bullet("**The equivalence is load-bearing on one assumption.** Dropping the "
       "fixed-prediction assumption changes that same gradient from "
       "\u221211.370010 to **\u22123.413946** \u2014 a factor of three, not a "
       "rounding error. The agreement with BP is not robust; it is contingent.")
bullet("**Error propagates as a wavefront whose speed is set by graph "
       "topology.** The number of inference steps before a node first receives a "
       "non-zero error equals its *shortest path to the output*, exactly. For the "
       "chain I get arrival times [5,4,3,2,1,0]; adding explicit skips gives "
       "[3,2,2,1,1,0]. The consequence is a floor on the inference budget: "
       "**T \u2265 L** for explicit skips versus **T \u2265 2L** for folded ones. "
       "Skip connections are not a detail in PC \u2014 they halve the required "
       "relaxation.")
para("So if PC is worth pursuing here, it is not for gradient quality. The "
     "question I actually care about is whether **the relaxation process produces "
     "different (and for patch prediction, better) latent representations** than a "
     "forward pass does. That is a question about representations, not about "
     "credit assignment, and as far as I can tell it is open.")

# ------------------------------------------------------------------ 2
heading("2. Scaling the architecture: reproducing \u00b5PC")
para("To get to anything JEPA-shaped I need depth, residual connections, and "
     "autodiff. I therefore worked through \u00b5PC (Innocenti et al.), which "
     "reparameterises PC networks in a depth-\u00b5P style so that deep residual "
     "PCNs train at all. I reproduced their setting from scratch: a **30-layer "
     "fully-connected ResNet, width 128, on MNIST**, using the jpc library, with a "
     "standard-parameterisation (SP) control at identical depth, weights and "
     "learning rates.")

doc.add_picture("fig_mupc_results.png", width=Inches(6.7))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

para("**Figure 1.** All three panels are from my own runs at depth 30, width 128, "
     "batch 64, activity lr 0.5, weight lr 0.1 (Adam), seed 4329. (a) Mean "
     "activity norm per layer after a forward pass at initialisation; log scale; "
     "the output layer is omitted since its scaling differs by construction. "
     "(b) SP training loss per iteration until it becomes non-finite. "
     "(c) \u00b5PC test accuracy over 900 iterations (\u22481 epoch), full 10k "
     "test set.", size=9, color=GREY, space_after=10)

para("The control is unambiguous. Under SP the forward pass grows **8.5\u00d7** "
     "from the first to the last hidden layer at initialisation "
     "(6.53 \u2192 55.55), and training diverges within six iterations \u2014 the "
     "loss goes 36.5, 1.0\u00d710^16^, 2.1\u00d710^23^, 1.9\u00d710^28^, "
     "1.7\u00d710^33^, inf. Under \u00b5PC the same forward pass grows "
     "**1.23\u00d7** (11.17 \u2192 13.73) and the network trains smoothly to "
     "**93.61%** test accuracy in a single epoch. This is the paper\u2019s "
     "headline claim, reproduced independently.")

# ------------------------------------------------------------------ 3
heading("3. The result that actually surprised me")
para("I instrumented the training loop to record the PC energy *F(z)* at every "
     "one of the T = 30 inference steps, so I could check whether the relaxation "
     "was doing any work. It is doing almost none:")

table([["training iteration", "F(z) before inference", "F(z) after T=30 steps",
        "relative drop"],
       ["100", "0.1618", "0.1610", "0.49%"],
       ["200", "0.1228", "0.1222", "0.49%"],
       ["300", "0.0978", "0.0973", "0.51%"],
       ["400", "0.0935", "0.0930", "0.53%"]],
      [1.5, 1.6, 1.7, 1.1])
doc.add_paragraph().paragraph_format.space_after = Pt(2)

para("Thirty steps of gradient descent on the activities move the energy by half "
     "a percent, and the figure is stable across training. In other words, under "
     "\u00b5PC **the feedforward initialisation already lands essentially at the "
     "inference fixed point**. That is a coherent reading of why \u00b5PC works "
     "\u2014 the paper shows the inference landscape becomes badly ill-conditioned "
     "with depth, and the reparameterisation appears to sidestep that by making "
     "relaxation nearly unnecessary rather than by making it converge faster.")
para("But it is awkward for my project. If inference is a 0.5% correction to a "
     "forward pass, then the latents PC produces are, to a very good "
     "approximation, the latents a forward pass produces \u2014 and the "
     "\u201cricher representations\u201d hypothesis from \u00a71 has little room "
     "to be true in this regime. Either the effect I am looking for lives "
     "somewhere other than the discriminative-MNIST setting, or it needs a network "
     "where relaxation genuinely moves the activities. **This is the single point "
     "I would most like a second opinion on.**")
para("I should flag a limitation honestly: the depth sweep across {8, 16, 30, 64} "
     "that would show how this scales is written but has not finished running, and "
     "my current numbers are single-seed. I am not claiming statistical strength "
     "for the 0.5% figure \u2014 only that it is consistent across four "
     "widely-separated logging points in one run.", size=9.2, color=GREY)

# ------------------------------------------------------------------ 4
heading("4. A direction I have just been pointed to: MPC / NeuroSSL")
para("Prof. Alexander Ororbia, whom I contacted about this project, pointed me to "
     "**Meta-Representational Predictive Coding** (Ororbia, Friston and Rao, "
     "arXiv:2503.21796), which he describes as the first encoder-centric or "
     "joint-embedding form of predictive coding. From its abstract, MPC avoids "
     "learning a generative model of raw pixels by predicting *representations* of "
     "the input across parallel streams, giving an encoder-only learning and "
     "inference scheme, with representational dynamics driven by active sensory "
     "glimpsing. He also noted that MPC ties the feedback weights to the forward "
     "ones \u2014 the usual PC simplification \u2014 and that there is no reason a "
     "structure like mine could not instead **learn the error weights**, as in "
     "neural generative coding (NGC).")
para("**I have not read this paper yet**, so I am not going to characterise its "
     "results here; it is the next thing I read, and I am flagging it because it "
     "clearly changes the shape of the project. Two implications I can already "
     "see:")
bullet("It suggests the joint-embedding + PC combination is viable in a form much "
       "closer to what I proposed than anything in the \u00b5PC line, which is "
       "purely discriminative.")
bullet("Learned (untied) error weights are a concrete axis I had not considered at "
       "all. Everything I have built so far assumes tied feedback, which is exactly "
       "the assumption that made my BP-equivalence check come out at cosine 1.0. "
       "Untying it would break that equivalence \u2014 which, given \u00a71, might "
       "be the point rather than a problem.")

# ------------------------------------------------------------------ 5
heading("5. What I have running, and what is next")
table([["component", "status"],
       ["PC over arbitrary computation graphs (JAX), 3 topologies",
        "done, verified against autodiff"],
       ["Error-wavefront / inference-budget analysis",
        "done; T \u2265 L vs T \u2265 2L result is exact"],
       ["\u00b5PC 30-layer ResNet on MNIST + SP control",
        "done, reproduces the paper"],
       ["Per-step inference-energy instrumentation",
        "done; gives the 0.5% result above"],
       ["Depth sweep {8, 16, 30, 64}, multi-seed", "written, not yet run"],
       ["Swap MNIST classification for patch prediction", "next"],
       ["Read MPC; evaluate untied / learned error weights", "next"]],
      [3.6, 3.1])
doc.add_paragraph().paragraph_format.space_after = Pt(2)

para("The immediate plan is to keep the \u00b5PC-parameterised residual encoder "
     "but replace the supervised objective with a JEPA-style one \u2014 predict "
     "the representation of a masked patch from the representation of the context "
     "\u2014 and then re-run the same energy instrumentation. If relaxation still "
     "only moves the energy half a percent in that setting, the representational "
     "hypothesis is in trouble and I would rather find that out in the next two "
     "weeks than in six months.")

# ------------------------------------------------------------------ 6
heading("6. Questions I would like your opinion on")
para("**1. Is the \u201cricher latents\u201d framing the right question, or am I "
     "chasing something that the 0.5% energy result has already answered "
     "negatively?** If the forward pass lands at the fixed point, is there a "
     "regime \u2014 different objective, weaker skips, smaller inference learning "
     "rate, untied weights \u2014 where relaxation would genuinely reshape the "
     "representation, or is that simply not what PC buys you?", indent=0.22)
para("**2. How would you actually measure \u201cricher\u201d?** This is my biggest "
     "methodological gap. Candidates I have considered: linear-probe accuracy on "
     "held-out tasks, effective rank of the representation matrix, resistance to "
     "representational collapse, downstream patch-prediction error. None feels "
     "decisive, and the \u00b5PC paper explicitly does not study representation "
     "quality at all.", indent=0.22)
para("**3. Are there directions in this space I should know about?** The MPC "
     "pointer above reoriented the project in a single email; I would rather "
     "discover the rest of that literature now than after building the wrong "
     "thing. In particular, is anyone applying NGC-style learned error weights in "
     "a joint-embedding setting, and is there work on PC for masked prediction "
     "specifically?", indent=0.22)
para("And the blunt version: **are these efforts pointed in a sensible "
     "direction?** I would rather be redirected early. All code, notebooks and "
     "figures are available if any of the above is worth looking at directly.")

para("References: Millidge, Tschantz & Buckley, *Predictive Coding Approximates "
     "Backprop along Arbitrary Computation Graphs* (2020). Innocenti et al., "
     "*\u00b5PC: Scaling Predictive Coding to 100+ Layer Networks*. Ororbia, "
     "Friston & Rao, *Meta-Representational Predictive Coding*, arXiv:2503.21796. "
     "Assran et al., *I-JEPA* (2023). Implementation uses the jpc library (Buckley "
     "lab).", size=9, color=GREY)

doc.save("pc_progress_report.docx")
print("saved")
